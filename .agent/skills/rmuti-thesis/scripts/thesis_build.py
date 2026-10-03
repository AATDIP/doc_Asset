"""Build one thesis chapter from input/chapters/chN.txt, formatted like the
senior thesis that passed review (input/reference/ไฟล์เล่มรุ่นพี่ที่ตรวจผ่านแล้ว.docx).

Usage (from repo root):  python .agent/skills/rmuti-thesis/scripts/thesis_build.py 2 [--no-pdf]

The document starts from input/reference/rmuti_thesis_template.docx, so fonts,
Heading 1-4 styles and their auto-numbering ("บทที่ %1", "%1.%2", ...) are the
template's own. Chapter file format (UTF-8, one paragraph per line):
    line 1         chapter number, e.g. "2"
    line 2         chapter title
    "2.1 ข้อความ"   -> Heading 2 (number is stripped, Word numbers it)
    "2.1.1 ข้อความ" -> chapter 1: Heading 3 (auto number)
                     chapter 2+: Normal, typed number, 1 tab   (as in senior book)
    "2.1.1.1 ..."   -> chapter 1: Heading 4 · chapter 2+: Normal, 2 tabs
    "@toc"          -> list of all X.X topics (senior chapter 2 intro)
    "1) ข้อความ"     -> Normal, 3 tabs (senior's list under X.X.X.X)
    "@fig ชื่อรูป"   -> framed 1x1 box with a placeholder + caption "รูปที่ N.n  ชื่อรูป" inside
    "@table ชื่อ"    -> caption "ตารางที่ N.n  ชื่อ" (left) + Table Grid built from the
                       following "a | b | c" lines; first row is the header
    other text      -> Normal body, 1 tab (2 tabs when under an X.X.X item)
"""
import os
import re
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", ".."))
TEMPLATE = os.path.join(ROOT, "input", "reference", "rmuti_thesis_template.docx")
NUM_RE = re.compile(r"^(\d+(?:\.\d+)+)\s+")
LIST_RE = re.compile(r"^\d+\)\s")
# column widths (cm) by column count; text width is 14.65 cm
COL_CM = {2: [4.6, 10.05], 3: [2.2, 5.6, 6.85], 4: [3.6, 2.9, 4.7, 3.45]}


def add_caption(container, label, chapter, text, align):
    """Caption paragraph "<label> N.n  text" with a SEQ field, so Word's
    Table of Figures can list it. The number is cached so it shows before F9."""
    add = container.add_paragraph
    p = add(style="Caption")
    p.alignment = align
    p.add_run(f"{label} {chapter}.")
    seq_no = container._seq[label] = container._seq.get(label, 0) + 1
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" SEQ {label} \\* ARABIC ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = str(seq_no)
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    p.add_run(f"  {text}")
    return p


def add_figure(doc, chapter, title):
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    trPr = box.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))  # picture and caption stay on one page
    cell = box.rows[0].cells[0]
    ph = cell.paragraphs[0]
    ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ph.paragraph_format.keep_with_next = True
    ph.add_run(f"[ใส่รูป: {title}]")
    cell._seq = doc._seq
    add_caption(cell, "รูปที่", chapter, title, WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()


def add_data_table(doc, chapter, title, rows):
    add_caption(doc, "ตารางที่", chapter, title, WD_ALIGN_PARAGRAPH.LEFT)
    cells = [[c.strip() for c in r.split("|")] for r in rows]
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    widths = COL_CM.get(ncol)
    for i, row in enumerate(cells):
        for j in range(ncol):
            cell = t.rows[i].cells[j]
            if widths:
                cell.width = Cm(widths[j])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(row[j] if j < len(row) else "")
    trPr = t.rows[0]._tr.get_or_add_trPr()  # repeat header row on each page
    trPr.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph()


def mark_thai(run):
    """Tag a run as Thai complex script. Without <w:cs/> Word treats Thai as
    Latin, can't word-break it, and thaiDistribute spreads the letters apart."""
    rPr = run._r.get_or_add_rPr()
    for tag in ("w:rFonts", "w:cs", "w:lang"):
        for old in rPr.findall(qn(tag)):
            rPr.remove(old)
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(a), "TH SarabunPSK")
    fonts.set(qn("w:hint"), "cs")
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "en-US")
    lang.set(qn("w:bidi"), "th-TH")
    rPr.insert(0, fonts)
    rPr.append(OxmlElement("w:cs"))
    rPr.append(lang)


def start_chapter_numbering(doc, chapter):
    """Heading styles share one numbering list; make it start at this chapter."""
    num_id = doc.styles["Heading 1"].element.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
    numbering = doc.part.numbering_part.element
    num = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == num_id)
    for old in num.findall(qn("w:lvlOverride")):
        num.remove(old)
    ov = OxmlElement("w:lvlOverride")
    ov.set(qn("w:ilvl"), "0")
    so = OxmlElement("w:startOverride")
    so.set(qn("w:val"), str(chapter))
    ov.append(so)
    num.append(ov)


def build_doc(lines):
    lines = [l.rstrip() for l in lines if l.strip()]
    chapter, title, body = int(lines[0]), lines[1].strip(), lines[2:]
    doc = docx.Document(TEMPLATE)
    for el in list(doc.element.body):  # keep only the template's section settings
        if el.tag != qn("w:sectPr"):
            doc.element.body.remove(el)
    s = doc.sections[0]
    # ponytail: one section at 3.81 cm; first page reaches 5.08 cm via 36pt
    # space-before on Heading 1 (senior book used a manual section break instead).
    s.top_margin = Cm(3.81)
    start_chapter_numbering(doc, chapter)

    h1 = doc.add_paragraph(style="Heading 1")
    h1.paragraph_format.space_before = Pt(36)
    h1.add_run().add_break()  # "บทที่ N" (auto) / line break / title
    h1.add_run(title)

    topics = [NUM_RE.sub("", l) for l in body if (m := NUM_RE.match(l)) and m.group(1).count(".") == 1]
    merged, just_merged = [], False
    for line in (l.strip() for l in body):
        prev = merged[-1] if merged else ""
        pm = NUM_RE.match(prev)
        if (chapter > 1 and not just_merged and pm and pm.group(1).count(".") >= 2
                and not NUM_RE.match(line) and not line.startswith("@")
                and " | " not in line and not LIST_RE.match(line)
                and not prev.endswith("ดังนี้")):
            merged[-1] = f"{prev} {line}"
            just_merged = True  # only the first body line joins the item
        else:
            merged.append(line)
            just_merged = False

    doc._seq = {}
    first_h2, under_item = True, False
    i = 0
    while i < len(merged):
        line = merged[i]
        i += 1
        if line.startswith("@fig "):
            add_figure(doc, chapter, line[5:].strip())
            continue
        if line.startswith("@table "):
            rows = []
            while i < len(merged) and " | " in merged[i]:
                rows.append(merged[i])
                i += 1
            add_data_table(doc, chapter, line[7:].strip(), rows)
            continue
        if LIST_RE.match(line):
            doc.add_paragraph("			" + line)
            continue
        if line == "@toc":
            for n, t in enumerate(topics, 1):
                p = doc.add_paragraph(f"{chapter}.{n} {t}")
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Cm(1.27)
            continue
        m = NUM_RE.match(line)
        level = m.group(1).count(".") + 1 if m else 0
        if level == 2:
            if first_h2 and chapter == 1:
                doc.add_paragraph()  # senior: blank line under the chapter title
            elif not first_h2:
                doc.add_paragraph()  # senior: blank line before each X.X
            first_h2, under_item = False, False
            h = doc.add_paragraph(NUM_RE.sub("", line), style="Heading 2")
            h.paragraph_format.keep_with_next = True  # no heading stranded at page bottom
        elif level >= 3 and chapter == 1:
            doc.add_paragraph(NUM_RE.sub("", line), style=f"Heading {min(level, 4)}")
        elif level >= 3:
            p = doc.add_paragraph("\t" * (level - 2) + line)
            # label-only item ("2.1.1 ข้อดี") stays with the list under it
            p.paragraph_format.keep_with_next = " " not in NUM_RE.sub("", line)
            under_item = True
        else:
            doc.add_paragraph(("\t\t" if under_item else "\t") + line)
    for p in all_paragraphs(doc):
        for r in p.runs:
            mark_thai(r)
    return doc, chapter, title


def all_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def check(doc):
    """Fail loudly if output drifts from the senior-book format."""
    paras = doc.paragraphs
    assert paras[0].style.name == "Heading 1"
    for p in list(all_paragraphs(doc))[1:]:
        if p.style.name == "Heading 2":
            assert not NUM_RE.match(p.text), f"typed number on Heading 2: {p.text[:40]}"
        for r in p.runs:
            assert r._r.rPr.find(qn("w:cs")) is not None, f"run not tagged Thai: {p.text[:40]}"
            assert r.bold is None, f"direct bold (use styles): {p.text[:40]}"
        assert p.paragraph_format.first_line_indent is None, f"use tabs, not indent: {p.text[:40]}"


FONT_FILE = "C:/Windows/Fonts/THSarabunNew.ttf"  # same metrics as TH SarabunPSK
LINE_PT = 415.0  # A4 21 cm - 3.81 - 2.54 margins = 14.65 cm
TAB_PT = 36.0    # template default tab stop 1.27 cm
# where the text starts on line 1 (pt); from the template's numbering indents
FIRST_OFFSET_PT = {"Heading 3": 58.5, "Heading 4": 42.6}


def short_last_lines(doc, min_fill=0.4):
    """Estimate wrapping and list paragraphs whose last line holds only a
    word or two (looks ragged; advisors ask to fix). Approximate: sums glyph
    advances, ignores Word's exact Thai break points."""
    try:
        from PIL import ImageFont
        font = ImageFont.truetype(FONT_FILE, 16)
    except (ImportError, OSError):
        return []
    bad = []
    for p in doc.paragraphs:
        text = p.text.replace("	", "")
        if not text.strip() or p.style.name in ("Heading 1", "Heading 2", "Caption"):
            continue
        width = font.getlength(text)
        if p.style.name in FIRST_OFFSET_PT:
            level = int(p.style.name[-1])
            width += font.getlength(".".join(["0"] * level) + "  ")
            first = LINE_PT - FIRST_OFFSET_PT[p.style.name]
        else:
            first = LINE_PT - TAB_PT * p.text.count("	", 0, 3)
        # >95% full is risky too: estimate error can push one word to a new line
        if width <= first:
            if width > 0.95 * first:
                bad.append((round(width / first * 100), text))
            continue
        last = (width - first) % LINE_PT
        if last < min_fill * LINE_PT or last > 0.95 * LINE_PT:
            bad.append((round(last / LINE_PT * 100), text))
    return bad


def export_pdf(docx_path, pdf_path):
    import win32com.client
    word = win32com.client.DispatchEx("Word.Application")  # own instance; Dispatch can grab one that is still quitting
    word.Visible = False
    try:
        d = word.Documents.Open(docx_path)
        d.SaveAs(pdf_path, FileFormat=17)
        d.Close()
    finally:
        word.Quit()


def render_pngs(pdf_path, stem):
    import pymupdf
    out_dir = os.path.join(ROOT, "process", "temp")
    os.makedirs(out_dir, exist_ok=True)
    pdf = pymupdf.open(pdf_path)
    for i, page in enumerate(pdf):
        page.get_pixmap(dpi=110).save(os.path.join(out_dir, f"{stem}_p{i + 1}.png"))
    return len(pdf)


def main(argv):
    n = argv[1]
    with open(os.path.join(ROOT, "input", "chapters", f"ch{n}.txt"), encoding="utf-8") as f:
        doc, chapter, title = build_doc(f.readlines())
    check(doc)
    for fill, text in short_last_lines(doc):
        print(f"WARN last line ~{fill}% full, reword: {text[:50]}")
    stem = f"Asset_บทที่{chapter}_{title.replace(' ', '')}"
    out = os.path.join(ROOT, "output", stem + ".docx")
    try:
        doc.save(out)
    except PermissionError:  # file open in Word: don't lose the build
        out = os.path.join(ROOT, "output", stem + "_new.docx")
        doc.save(out)
        print("Target open in Word, saved instead to", out)
    print("docx:", out)
    if "--no-pdf" in argv:
        return
    pdf = os.path.join(ROOT, "output", stem + ".pdf")
    export_pdf(out, pdf)
    print(f"pdf: {pdf} ({render_pngs(pdf, stem)} pages, previews in process/temp/)")


if __name__ == "__main__":
    main(sys.argv)
